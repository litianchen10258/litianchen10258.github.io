import os, re, sys, json, shutil, requests, html
from datetime import datetime
from pathlib import Path

def log(msg):
    with open(Path(__file__).parent / '爬取日志.txt', 'a', encoding='utf-8') as f:
        f.write(f"[{datetime.now().strftime('%Y-%m-%d %H:%M:%S')}] {msg}\n")
from urllib.parse import urlparse, unquote

def extract_novel_id(url):
    m = re.search(r'pixiv\.net/.*?novel/.*?(\d+)', url)
    if not m:
        print('错误：无法从URL提取小说ID')
        sys.exit(1)
    return m.group(1)

def main():
    batch_mode = '--batch' in sys.argv

    if batch_mode:
        url_file = Path(__file__).parent / 'pixiv_urls.txt'
        if not url_file.exists():
            print(f'批量模式需要创建 {url_file}，每行一个小说URL')
            sys.exit(1)
        urls = [l.strip() for l in url_file.read_text(encoding='utf-8').splitlines() if l.strip()]
        if not urls:
            print('pixiv_urls.txt 为空')
            sys.exit(1)
    elif len(sys.argv) < 2:
        print('用法: python pixiv_novel_saver.py <novel_url>')
        print('       python pixiv_novel_saver.py --batch')
        sys.exit(1)
    else:
        urls = [sys.argv[1]]

    cookie_file = Path(__file__).parent / 'pixiv_cookies.txt'
    if not cookie_file.exists():
        print(f'请创建 {cookie_file}，放入你的Pixiv登录cookie（格式 PHPSESSID=xxx; 每行一个）')
        sys.exit(1)

    cookies = {}
    for line in cookie_file.read_text(encoding='utf-8').strip().splitlines():
        line = line.strip()
        if '=' in line:
            k, v = line.split('=', 1)
            cookies[k.strip()] = v.strip()

    session = requests.Session()
    session.cookies.update(cookies)
    session.headers.update({'User-Agent': 'Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36'})

    for url in urls:
        novel_id = extract_novel_id(url)
        print(f"\n--- 处理: {url} ---")
        _download_novel(session, novel_id)

def _download_novel(session, novel_id):
    log(f'小说 {novel_id} 开始')
    print('获取小说信息...')
    info_url = f'https://www.pixiv.net/ajax/novel/{novel_id}'
    resp = session.get(info_url)
    info = resp.json()
    if info.get('error'):
        print(f"获取失败 (HTTP {resp.status_code}): {info.get('message', '')}")
        print('可能原因：① cookie 过期或无效 ② 小说ID错误 ③ 需要年龄验证')
        log(f'小说 {novel_id} 失败: HTTP {resp.status_code}')
        return

    body = info['body']
    title = body.get('title', f'novel_{novel_id}')
    series_title = body.get('seriesTitle', '')
    cover_url = body.get('coverUrl', '')
    tags = [t.get('tag', '') for t in body.get('tags', {}).get('tags', [])]
    description = html.unescape(re.sub(r'<[^>]+>', '', body.get('description', '')))

    print('获取小说正文...')
    text_url = f'https://www.pixiv.net/ajax/novel/{novel_id}/text'
    text_data = session.get(text_url).json()
    if text_data.get('error') and isinstance(text_data['body'], list) and len(text_data['body']) == 0:
        # 回退到 info API 中的 content 字段
        raw_html = body.get('content', '')
    else:
        raw_html = ''.join(text_data['body']) if isinstance(text_data['body'], list) else text_data['body']

    out_dir = Path(__file__).parent / f'novel_{novel_id}'
    out_dir.mkdir(exist_ok=True)
    img_dir = out_dir / 'images'
    img_dir.mkdir(exist_ok=True)

    images_map = {}

    if cover_url:
        print('下载封面...')
        cover_ext = os.path.splitext(cover_url.split('?')[0])[1] or '.jpg'
        r = session.get(cover_url, headers={'Referer': 'https://www.pixiv.net/'})
        if r.ok:
            cover_path = img_dir / f'cover{cover_ext}'
            cover_path.write_bytes(r.content)
            images_map[cover_url] = f'images/cover{cover_ext}'

    print('下载插图...')
    # 途径1: textEmbeddedImages
    tei = body.get('textEmbeddedImages') or {}
    for img_id, img_info in tei.items():
        img_url = img_info.get('url') or img_info.get('urls', {}).get('original', '')
        if not img_url:
            continue
        ext = os.path.splitext(img_url.split('?')[0])[1] or '.jpg'
        local_name = f'img_{img_id}{ext}'
        r = session.get(img_url, headers={'Referer': 'https://www.pixiv.net/'})
        if r.ok:
            (img_dir / local_name).write_bytes(r.content)
            images_map[f'[uploadedimage:{img_id}]'] = f'images/{local_name}'
            print(f'  下载: {local_name}')

    # 途径2: pageList
    for page in (body.get('pageList') or []):
        img_id = page.get('id', '')
        urls = page.get('urls', {})
        img_url = urls.get('original', urls.get('thumb', ''))
        if not img_id or not img_url:
            continue
        ext = os.path.splitext(img_url.split('?')[0])[1] or '.jpg'
        local_name = f'img_{img_id}{ext}'
        r = session.get(img_url, headers={'Referer': 'https://www.pixiv.net/'})
        if r.ok:
            (img_dir / local_name).write_bytes(r.content)
            images_map[f'[uploadedimage:{img_id}]'] = f'images/{local_name}'
            print(f'  下载: {local_name}')

    # 途径3: 从正文提取插图ID并猜URL
    for m in re.finditer(r'\[uploadedimage:(\d+)\]', raw_html):
        img_id = m.group(1)
        key = f'[uploadedimage:{img_id}]'
        if key in images_map:
            continue
        found = False
        for try_url in [
            f'https://i.pximg.net/novel-upload-original/img/{img_id}_master1200.jpg',
            f'https://i.pximg.net/novel-upload-original/img/{img_id}.jpg',
            f'https://i.pximg.net/img-original/img/0000/00/00/00/00/00/{img_id}_p0.png',
        ]:
            r = session.get(try_url, headers={'Referer': 'https://www.pixiv.net/'})
            if r.ok:
                ext = os.path.splitext(try_url.split('?')[0])[1] or '.jpg'
                local_name = f'img_{img_id}{ext}'
                (img_dir / local_name).write_bytes(r.content)
                images_map[key] = f'images/{local_name}'
                print(f'  下载: {local_name}')
                found = True
                break
        if not found:
            images_map[key] = None

        print('生成Word...')
    from docx import Document
    from docx.shared import Inches, Pt
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    doc = Document()
    style = doc.styles['Normal']
    style.font.name = 'Yu Mincho'
    style.font.size = Pt(11)
    style.paragraph_format.line_spacing = 1.5

    doc.add_heading(html.unescape(title), level=1)

    meta_tags_str = ', '.join(tags[:10])
    if series_title:
        doc.add_paragraph('系列: ' + html.unescape(series_title))

    if tags:
        doc.add_paragraph('标签: ' + meta_tags_str)

    if description:
        doc.add_paragraph(html.unescape(description))

    cover_path = img_dir / 'cover.jpg'
    if cover_path.exists():
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(str(cover_path), width=Inches(3))

    doc.add_paragraph('_' * 40)

    for line in raw_html.split(chr(10)):
        line = line.strip()
        if not line:
            continue
        img_keys = re.findall(r'\[uploadedimage:(\d+)\]', line)
        if img_keys:
            for key in img_keys:
                local = images_map.get('[uploadedimage:' + key + ']')
                if local:
                    img_path = img_dir / os.path.basename(local)
                    if img_path.exists():
                        p = doc.add_paragraph()
                        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                        p.add_run().add_picture(str(img_path), width=Inches(5))
                else:
                    doc.add_paragraph('[插图 ' + key + ' 下载失败]')
        else:
            doc.add_paragraph(html.unescape(line))

    out_file = out_dir.parent / (re.sub(r'[\\/:*?"<>|]', '', title).strip() + '.docx')
    doc.save(str(out_file))
    log(f'小说 {novel_id} 《{title}》成功 -> {out_file}')
    print('完成! 保存至: ' + str(out_file))
    shutil.rmtree(out_dir, ignore_errors=True)
if __name__ == '__main__':
    main()
